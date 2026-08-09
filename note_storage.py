from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


SITE_NOTES_FILE = Path(__file__).with_name("Site_Notes.xlsx")
SHEET_NAME = "Site Notes"


def ensure_notes_workbook():
    """Create the notes workbook if it does not already exist."""

    if SITE_NOTES_FILE.exists():
        return

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    sheet.append(
        [
            "UPRN",
            "Date Time",
            "Site Access Note",
        ]
    )

    workbook.save(SITE_NOTES_FILE)
    workbook.close()


def save_site_note(uprn: str, note: str):
    """Append a site access note to the Excel workbook."""

    ensure_notes_workbook()

    workbook = load_workbook(SITE_NOTES_FILE)
    sheet = workbook[SHEET_NAME]

    sheet.append(
        [
            uprn.strip().upper(),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            note.strip(),
        ]
    )

    workbook.save(SITE_NOTES_FILE)
    workbook.close()


def load_site_notes(uprn: str) -> list[dict]:
    """Return all saved notes for a UPRN, newest first."""

    ensure_notes_workbook()

    workbook = load_workbook(
        SITE_NOTES_FILE,
        data_only=True,
    )

    sheet = workbook[SHEET_NAME]

    requested_uprn = uprn.strip().upper()
    notes = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        saved_uprn, date_time, note = row

        if saved_uprn is None:
            continue

        if str(saved_uprn).strip().upper() == requested_uprn:
            notes.append(
                {
                    "UPRN": saved_uprn,
                    "Date Time": date_time,
                    "Site Access Note": note,
                }
            )

    workbook.close()

    return list(reversed(notes))
